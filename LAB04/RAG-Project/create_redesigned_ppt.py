import pptx
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = pptx.Presentation()

# ----------------- SLIDE 1 -----------------
slide_layout = prs.slide_layouts[5] # blank slide with title
slide = prs.slides.add_slide(slide_layout)
title_shape = slide.shapes.title
title_shape.text = "RAG System Architecture & Workflow"
title_shape.text_frame.paragraphs[0].font.size = Pt(36)
title_shape.text_frame.paragraphs[0].font.bold = True

def add_clean_box(slide, left, top, width, height, title, subtitle, fill_color, border_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 0, 0)
    p.alignment = PP_ALIGN.CENTER
    
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(80, 80, 80)
        p2.alignment = PP_ALIGN.CENTER
    return shape

def add_clean_arrow(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(220, 220, 220)
    shape.line.color.rgb = RGBColor(180, 180, 180)
    return shape

def add_section_label(slide, left, top, width, height, text):
    tf = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    p = tf.text_frame.paragraphs[0]
    p.text = text
    p.font.bold = True
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(50, 50, 50)


# 1. BUILD
add_section_label(slide, 0.2, 1.2, 2.0, 0.5, "1. BUILD")
c_build = RGBColor(237, 245, 255) # Light Blue
b_build = RGBColor(100, 150, 255)
y1 = 1.6
add_clean_box(slide, 0.4, y1, 1.7, 0.9, "Knowledge Base", "data/qa_looseweight.txt", c_build, b_build)
add_clean_arrow(slide, 2.15, y1+0.35, 0.2, 0.2)
add_clean_box(slide, 2.4, y1, 1.7, 0.9, "document_loader", "Parse Q&A", c_build, b_build)
add_clean_arrow(slide, 4.15, y1+0.35, 0.2, 0.2)
add_clean_box(slide, 4.4, y1, 1.7, 0.9, "text_splitter", "Chunking", c_build, b_build)
add_clean_arrow(slide, 6.15, y1+0.35, 0.2, 0.2)
add_clean_box(slide, 6.4, y1, 1.7, 0.9, "embedding_model", "Vectorization", c_build, b_build)
add_clean_arrow(slide, 8.15, y1+0.35, 0.2, 0.2)
add_clean_box(slide, 8.4, y1-0.2, 1.4, 1.3, "vector_db", "FAISS Index\nBM25 Index\nChunk Store", RGBColor(243, 229, 245), RGBColor(156, 39, 176))

# 2. QUERY
add_section_label(slide, 0.2, 3.1, 2.0, 0.5, "2. QUERY")
c_query = RGBColor(240, 248, 235) # Light Green
b_query = RGBColor(100, 180, 100)
c_orange = RGBColor(255, 245, 235)
b_orange = RGBColor(255, 150, 50)
y2 = 3.5

add_clean_box(slide, 0.4, y2, 1.5, 0.9, "Input", "query_transform.py\n(Normalize)", c_query, b_query)
add_clean_arrow(slide, 1.95, y2+0.35, 0.2, 0.2)
add_clean_box(slide, 2.2, y2, 1.6, 0.9, "Retrieval", "hybrid_retriever.py\n(FAISS + BM25)", c_orange, b_orange)
add_clean_arrow(slide, 3.85, y2+0.35, 0.2, 0.2)
add_clean_box(slide, 4.1, y2, 1.6, 0.9, "Context", "prompt_templates.py\n(Format Blocks)", c_orange, b_orange)
add_clean_arrow(slide, 5.75, y2+0.35, 0.2, 0.2)
add_clean_box(slide, 6.0, y2, 1.6, 0.9, "LLM", "generator.py\n(Ollama/OpenAI)", c_orange, b_orange)
add_clean_arrow(slide, 7.65, y2+0.35, 0.2, 0.2)
add_clean_box(slide, 7.9, y2, 1.5, 0.9, "Output", "Return Answer\n+ Citations", c_query, b_query)

# Memory box underneath query
add_clean_box(slide, 3.1, 4.6, 4.0, 0.5, "memory.py", "Maintains last 6 turns for follow-up context", RGBColor(250, 250, 250), RGBColor(200, 200, 200))


# 3. EVALUATION
add_section_label(slide, 0.2, 5.5, 3.0, 0.5, "3. EVALUATION")
c_eval = RGBColor(255, 240, 245) # Pinkish
b_eval = RGBColor(200, 100, 150)
y3 = 5.9

add_clean_box(slide, 0.4, y3, 2.0, 0.9, "Golden Set", "golden_set.json\n(Test Cases)", c_eval, b_eval)
add_clean_arrow(slide, 2.45, y3+0.35, 0.3, 0.2)
add_clean_box(slide, 2.8, y3, 2.5, 0.9, "Evaluate", "eval_retrieval.py (Hit@k)\neval_generation.py (Quality)", RGBColor(250, 250, 250), b_eval)
add_clean_arrow(slide, 5.35, y3+0.35, 0.3, 0.2)
add_clean_box(slide, 5.7, y3, 2.0, 0.9, "Report", "metrics.py\nCompare & Improve", c_eval, b_eval)



# ----------------- SLIDE 2 -----------------
slide_layout2 = prs.slide_layouts[6] # completely blank
slide2 = prs.slides.add_slide(slide_layout2)

# Custom big title
tf_title = slide2.shapes.add_textbox(Inches(0.4), Inches(0.4), Inches(9.2), Inches(0.8))
p_title = tf_title.text_frame.paragraphs[0]
p_title.text = "RAG Workflow Explanation"
p_title.font.bold = True
p_title.font.size = Pt(32)

# Two column layout for text
col1 = slide2.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(4.5), Inches(5.5))
tf1 = col1.text_frame
tf1.word_wrap = True

def add_stage(tf, stage, happens, component, result):
    # Stage Name
    p = tf.add_paragraph() if tf.text else tf.paragraphs[0]
    p.text = stage
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(0, 100, 200)
    p.space_before = Pt(10)
    
    # Details
    def add_line(label, text):
        pl = tf.add_paragraph()
        pl.text = f"→ {label}: {text}"
        pl.font.size = Pt(12)
        pl.font.color.rgb = RGBColor(50, 50, 50)
        
    add_line("What happens", happens)
    add_line("Component", component)
    add_line("Result", result)


add_stage(tf1, "1. Input",
          "User asks a question. Slang is normalized and queries can be rewritten.",
          "query_transform.py",
          "Optimized search query.")

add_stage(tf1, "2. Retrieval",
          "Searches the knowledge base using dense vectors and keywords.",
          "hybrid_retriever.py (FAISS + BM25)",
          "Top relevant data chunks.")

add_stage(tf1, "3. Context",
          "Formats the retrieved chunks into numbered blocks and applies conversational history.",
          "prompt_templates.py",
          "A strict, highly structured prompt.")

col2 = slide2.shapes.add_textbox(Inches(5.1), Inches(1.3), Inches(4.5), Inches(5.5))
tf2 = col2.text_frame
tf2.word_wrap = True

add_stage(tf2, "4. LLM",
          "The language model reads the context and answers the question strictly using references.",
          "generator.py (Ollama/OpenAI)",
          "Generated answer with inline citations [1].")

add_stage(tf2, "5. Output",
          "Appends standard disclaimers and prints the final formatted result.",
          "main.py",
          "User receives a verified, accurate response.")

# EVALUATION SECTION
pe = tf2.add_paragraph()
pe.text = "EVALUATION - Measure and Improve"
pe.font.bold = True
pe.font.size = Pt(16)
pe.font.color.rgb = RGBColor(200, 50, 150)
pe.space_before = Pt(15)

def add_eval_line(label, text):
    pl = tf2.add_paragraph()
    pl.text = f"• {label}: {text}"
    pl.font.size = Pt(12)
    pl.font.color.rgb = RGBColor(50, 50, 50)

add_eval_line("Measure", "Evaluate generated answers against the expected answers in golden_set.json.")
add_eval_line("Identify", "Find weaknesses in retrieval (eval_retrieval.py) or answer quality (eval_generation.py).")
add_eval_line("Improve", "Use metrics.py reports to adjust settings and improve the RAG system over time.")

prs.save('D:\\ATCS-CPE\\ATCS-CPE-main\\LAB03\\RAG-Project\\RAG_System_Workflow_Redesigned.pptx')
print("Successfully generated redesigned PPTX.")
