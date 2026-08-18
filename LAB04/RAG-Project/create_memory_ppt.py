import pptx
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = pptx.Presentation()

# ----------------- SLIDE 1 -----------------
slide_layout = prs.slide_layouts[5] # blank slide with title
slide = prs.slides.add_slide(slide_layout)
title_shape = slide.shapes.title
title_shape.text = "RAG System Architecture & Memory Workflow"
title_shape.text_frame.paragraphs[0].font.size = Pt(32)
title_shape.text_frame.paragraphs[0].font.bold = True

def add_clean_box(slide, left, top, width, height, title, subtitle, fill_color, border_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0, 0, 0)
    p.alignment = PP_ALIGN.CENTER
    
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(10)
        p2.font.color.rgb = RGBColor(80, 80, 80)
        p2.alignment = PP_ALIGN.CENTER
    return shape

def add_clean_arrow(slide, left, top, width, height, rotation=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(220, 220, 220)
    shape.line.color.rgb = RGBColor(150, 150, 150)
    shape.rotation = rotation
    return shape

def add_text_label(slide, left, top, width, height, text, font_size=10, bold=False, color=RGBColor(100, 100, 100)):
    tf = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    p = tf.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.CENTER

def add_section_label(slide, left, top, width, height, text):
    tf = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    p = tf.text_frame.paragraphs[0]
    p.text = text
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(50, 50, 50)


# 1. BUILD
add_section_label(slide, 0.2, 1.0, 2.0, 0.5, "1. BUILD")
c_build = RGBColor(237, 245, 255) # Light Blue
b_build = RGBColor(100, 150, 255)
y1 = 1.4
add_clean_box(slide, 0.4, y1, 1.6, 0.8, "Knowledge Base", "data/qa.txt", c_build, b_build)
add_clean_arrow(slide, 2.05, y1+0.3, 0.2, 0.15)
add_clean_box(slide, 2.3, y1, 1.6, 0.8, "document_loader", "Parse Q&A", c_build, b_build)
add_clean_arrow(slide, 3.95, y1+0.3, 0.2, 0.15)
add_clean_box(slide, 4.2, y1, 1.6, 0.8, "text_splitter", "Chunking", c_build, b_build)
add_clean_arrow(slide, 5.85, y1+0.3, 0.2, 0.15)
add_clean_box(slide, 6.1, y1, 1.6, 0.8, "embedding_model", "Vectorization", c_build, b_build)
add_clean_arrow(slide, 7.75, y1+0.3, 0.2, 0.15)
add_clean_box(slide, 8.0, y1, 1.6, 0.8, "vector_db", "FAISS / BM25", RGBColor(243, 229, 245), RGBColor(156, 39, 176))

# 2. QUERY
add_section_label(slide, 0.2, 2.7, 2.0, 0.5, "2. QUERY")
c_query = RGBColor(240, 248, 235) # Light Green
b_query = RGBColor(100, 180, 100)
c_orange = RGBColor(255, 245, 235)
b_orange = RGBColor(255, 150, 50)
y2 = 3.1

add_clean_box(slide, 0.4, y2, 1.5, 0.8, "Input", "User Question", c_query, b_query)
add_clean_arrow(slide, 1.95, y2+0.3, 0.2, 0.15)
add_clean_box(slide, 2.2, y2, 1.6, 0.8, "query_transform", "Normalize & Rewrite", c_orange, b_orange)
add_clean_arrow(slide, 3.85, y2+0.3, 0.2, 0.15)
add_clean_box(slide, 4.1, y2, 1.6, 0.8, "Retrieval", "hybrid_retriever", c_orange, b_orange)
add_clean_arrow(slide, 5.75, y2+0.3, 0.2, 0.15)
add_clean_box(slide, 6.0, y2, 1.6, 0.8, "Context & LLM", "generator.py", c_orange, b_orange)
add_clean_arrow(slide, 7.65, y2+0.3, 0.2, 0.15)
add_clean_box(slide, 7.9, y2, 1.5, 0.8, "Output", "Final Answer", c_query, b_query)

# --- MEMORY WORKFLOW HIGHLIGHT ---
add_section_label(slide, 0.2, 4.2, 2.0, 0.5, "Memory Flow")
c_mem = RGBColor(255, 250, 230)
b_mem = RGBColor(200, 180, 50)
add_clean_box(slide, 4.0, 4.6, 2.5, 0.7, "memory.py", "Keeps last 6 turns\n(Configurable)", c_mem, b_mem)

# Arrow 1: Output -> Memory (Save)
add_clean_arrow(slide, 6.6, 4.85, 1.3, 0.15, rotation=180)
add_text_label(slide, 6.5, 5.0, 1.5, 0.4, "1. Saves generated answer", font_size=9, color=RGBColor(50,50,50))

# Arrow 2: Memory -> query_transform (Rewrite)
add_clean_arrow(slide, 2.9, 4.5, 1.1, 0.15, rotation=220)
add_text_label(slide, 2.4, 4.8, 1.8, 0.4, "2. history for is_followup()", font_size=9, color=RGBColor(50,50,50))

# Arrow 3: Memory -> generator (Context)
add_clean_arrow(slide, 5.5, 4.2, 1.1, 0.15, rotation=310)
add_text_label(slide, 5.6, 4.4, 1.8, 0.4, "3. history for LLM prompt", font_size=9, color=RGBColor(50,50,50))


# 3. EVALUATION
add_section_label(slide, 0.2, 5.8, 3.0, 0.5, "3. EVALUATION")
c_eval = RGBColor(255, 240, 245) # Pinkish
b_eval = RGBColor(200, 100, 150)
y3 = 6.2

add_clean_box(slide, 0.4, y3, 1.8, 0.8, "Golden Set", "golden_set.json", c_eval, b_eval)
add_clean_arrow(slide, 2.25, y3+0.3, 0.3, 0.15)
add_clean_box(slide, 2.6, y3, 2.3, 0.8, "Evaluate", "eval_retrieval & eval_generation", RGBColor(250, 250, 250), b_eval)
add_clean_arrow(slide, 4.95, y3+0.3, 0.3, 0.15)
add_clean_box(slide, 5.3, y3, 1.8, 0.8, "Report", "metrics.py (Compare)", c_eval, b_eval)


# ----------------- SLIDE 2 -----------------
slide_layout2 = prs.slide_layouts[6] # blank
slide2 = prs.slides.add_slide(slide_layout2)

# Custom title
tf_title = slide2.shapes.add_textbox(Inches(0.4), Inches(0.4), Inches(9.2), Inches(0.8))
p_title = tf_title.text_frame.paragraphs[0]
p_title.text = "Understanding the memory.py Workflow"
p_title.font.bold = True
p_title.font.size = Pt(28)

# Explanations Box
col1 = slide2.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(9.0), Inches(5.5))
tf1 = col1.text_frame
tf1.word_wrap = True

def add_flow_explanation(tf, trigger, intermediate, action, destination, result):
    # Main Header
    p = tf.add_paragraph() if tf.text else tf.paragraphs[0]
    p.text = trigger
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(0, 100, 200)
    p.space_before = Pt(10)
    
    # Sub items
    p1 = tf.add_paragraph()
    p1.text = f"→ {intermediate}"
    p1.font.size = Pt(12)
    p1.font.bold = True
    
    p2 = tf.add_paragraph()
    p2.text = f"→ {action}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = RGBColor(80, 80, 80)
    
    p3 = tf.add_paragraph()
    p3.text = f"→ {destination}"
    p3.font.size = Pt(12)
    p3.font.bold = True
    
    p4 = tf.add_paragraph()
    p4.text = f"→ {result}"
    p4.font.size = Pt(12)
    p4.font.color.rgb = RGBColor(80, 80, 80)

add_flow_explanation(tf1, 
    "1. Saving the Conversation (End of query cycle)",
    "rag_pipeline.py generates the final answer",
    "Calls memory.add_user() and memory.add_assistant() to save the interaction",
    "memory.py",
    "Stores the new Q&A turn. It deletes the oldest messages if it exceeds 6 turns, but preserves the very first topic.")

add_flow_explanation(tf1, 
    "2. Query Rewriting (Start of next query cycle)",
    "User asks a new question in main.py",
    "rag_pipeline.py calls memory.is_followup() to check if the question relies on context (e.g. 'why did that happen?')",
    "query_transform.py",
    "If True, the transformer reads the memory history and rewrites the short question into a complete, standalone search query.")

add_flow_explanation(tf1, 
    "3. Generating Context-Aware Answers",
    "hybrid_retriever.py successfully retrieves data chunks",
    "rag_pipeline.py fetches the full history via memory.get_context()",
    "generator.py (prompt_templates.py)",
    "The history is securely appended to the LLM's system prompt alongside the retrieved chunks, ensuring the model's answer is context-aware.")

prs.save('D:\\ATCS-CPE\\ATCS-CPE-main\\LAB03\\RAG-Project\\RAG_System_Workflow_Memory_Enhanced.pptx')
print("Successfully generated memory enhanced PPTX.")
