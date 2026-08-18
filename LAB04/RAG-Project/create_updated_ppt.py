import collections.abc
import pptx
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = pptx.Presentation()

# Slide 1: Workflow Diagram
slide_layout = prs.slide_layouts[5] # blank slide with title
slide = prs.slides.add_slide(slide_layout)
shapes = slide.shapes
title_shape = shapes.title
title_shape.text = "RAG System Workflow Diagram"

def add_box(slide, left, top, width, height, text, fill_color, border_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(0, 0, 0)
    p.alignment = PP_ALIGN.CENTER
    return shape

def add_arrow(slide, left, top, width, height, rotation=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(200, 200, 200)
    shape.line.color.rgb = RGBColor(150, 150, 150)
    shape.rotation = rotation
    return shape

def add_section_box(slide, left, top, width, height, title):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(248, 249, 250)
    shape.line.color.rgb = RGBColor(206, 212, 218)
    shape.line.width = Pt(2)
    
    tf = slide.shapes.add_textbox(Inches(left+0.1), Inches(top+0.1), Inches(width-0.2), Inches(0.5))
    tf.text_frame.text = title
    tf.text_frame.paragraphs[0].font.bold = True
    tf.text_frame.paragraphs[0].font.size = Pt(12)

# --- SECTION 1: BUILD ---
add_section_box(slide, 0.2, 1.1, 9.6, 1.6, "1. BUILD — Prepare dataset (python build_index.py)")

c1 = RGBColor(227, 242, 253) # blueish
b1 = RGBColor(33, 150, 243)
add_box(slide, 0.4, 1.6, 1.3, 0.9, "Knowledge Base\ndata/qa_looseweight.txt", c1, b1)
add_arrow(slide, 1.8, 2.0, 0.3, 0.15)
add_box(slide, 2.2, 1.6, 1.3, 0.9, "document_loader\nParse Q&A file", RGBColor(255, 255, 255), b1)
add_arrow(slide, 3.6, 2.0, 0.3, 0.15)
add_box(slide, 4.0, 1.6, 1.3, 0.9, "text_splitter\nbuild_chunks()", RGBColor(255, 255, 255), b1)
add_arrow(slide, 5.4, 2.0, 0.3, 0.15)
add_box(slide, 5.8, 1.6, 1.3, 0.9, "embedding_model\nMiniLM multilingual", RGBColor(255, 255, 255), b1)
add_arrow(slide, 7.2, 2.0, 0.3, 0.15)
add_box(slide, 7.6, 1.5, 1.8, 1.1, "vector_db/\ndocument.index\nbm25_index.pkl\nchunk_store.json", RGBColor(243, 229, 245), RGBColor(156, 39, 176))


# --- SECTION 2: QUERY ---
add_section_box(slide, 0.2, 2.9, 9.6, 2.3, "2. QUERY — Answer questions (python main.py)")

c2 = RGBColor(232, 245, 233) # greenish
b2 = RGBColor(76, 175, 80)
c3 = RGBColor(255, 243, 224) # orange
b3 = RGBColor(255, 152, 0)
add_box(slide, 0.4, 3.4, 1.2, 0.7, "Input\nUser Question", c2, b2)
add_arrow(slide, 1.7, 3.7, 0.2, 0.15)
add_box(slide, 2.0, 3.4, 1.2, 0.7, "query_transform\nNormalize / HyDE", RGBColor(250, 250, 250), b2)
add_arrow(slide, 3.3, 3.7, 0.2, 0.15)
add_box(slide, 3.6, 3.4, 1.2, 0.7, "Retrieval\nFAISS + BM25", c3, b3)
add_arrow(slide, 4.9, 3.7, 0.2, 0.15)
add_box(slide, 5.2, 3.4, 1.2, 0.7, "Context\nprompt_templates", c3, b3)
add_arrow(slide, 6.5, 3.7, 0.2, 0.15)
add_box(slide, 6.8, 3.4, 1.2, 0.7, "LLM\ngenerator.py", c3, b3)
add_arrow(slide, 8.1, 3.7, 0.2, 0.15)
add_box(slide, 8.4, 3.4, 1.0, 0.7, "Output\nAnswer", c2, b2)

# Memory box
add_box(slide, 3.6, 4.3, 2.8, 0.6, "memory.py\nKeep conversation history", RGBColor(250, 250, 250), b2)
add_arrow(slide, 5.0, 4.1, 0.2, 0.15, rotation=270)


# --- SECTION 3: EVALUATION ---
add_section_box(slide, 0.2, 5.4, 9.6, 1.9, "3. EVALUATION — Measure and improve")
c4 = RGBColor(243, 229, 245) # light purple
b4 = RGBColor(103, 58, 183)

add_box(slide, 0.4, 5.9, 1.5, 0.9, "golden_set.json\nTest cases\n(from chunk_store)", c4, b4)
add_arrow(slide, 2.0, 6.3, 0.3, 0.15)

add_box(slide, 2.4, 5.8, 1.6, 0.5, "eval_retrieval.py\nHit@k / MRR", RGBColor(250, 250, 250), b4)
add_arrow(slide, 4.1, 6.0, 0.3, 0.15)
add_box(slide, 4.5, 5.8, 1.8, 0.5, "Retrieval Report\noutputs/eval_retrieval.json", RGBColor(250, 250, 250), b4)

add_box(slide, 2.4, 6.5, 1.6, 0.5, "eval_generation.py\nFaithfulness", RGBColor(250, 250, 250), b4)
add_arrow(slide, 4.1, 6.7, 0.3, 0.15)
add_box(slide, 4.5, 6.5, 1.8, 0.5, "Generation Report\noutputs/eval_generation.json", RGBColor(250, 250, 250), b4)

add_arrow(slide, 6.4, 6.3, 0.3, 0.15)
add_box(slide, 6.8, 6.1, 1.4, 0.6, "metrics.py\nCompute & Compare", c4, b4)


# Slide 2: Workflow Explanation
slide_layout2 = prs.slide_layouts[1] # Title and Content
slide2 = prs.slides.add_slide(slide_layout2)
title_shape2 = slide2.shapes.title
title_shape2.text = "Workflow Explanation"

body_shape = slide2.shapes.placeholders[1]
tf2 = body_shape.text_frame

def add_level_text(tf, text, level=0, bold=False, size=18):
    p = tf.add_paragraph() if tf.text else tf.paragraphs[0]
    p.text = text
    p.level = level
    p.font.bold = bold
    p.font.size = Pt(size)

add_level_text(tf2, "1. Input", level=0, bold=True, size=16)
add_level_text(tf2, "User inputs a question via the terminal (main.py).", level=1, size=14)
add_level_text(tf2, "query_transform.py normalizes slang and formats the query.", level=1, size=14)

add_level_text(tf2, "2. Retrieval", level=0, bold=True, size=16)
add_level_text(tf2, "hybrid_retriever.py searches FAISS (Dense) and BM25 (Keyword).", level=1, size=14)
add_level_text(tf2, "Results are merged and scored via Reciprocal Rank Fusion.", level=1, size=14)

add_level_text(tf2, "3. Context", level=0, bold=True, size=16)
add_level_text(tf2, "prompt_templates.py formats top chunks into [1], [2] citation blocks.", level=1, size=14)

add_level_text(tf2, "4. LLM & 5. Output", level=0, bold=True, size=16)
add_level_text(tf2, "generator.py sends strict context and history to the LLM (Ollama/OpenAI/Gemini).", level=1, size=14)
add_level_text(tf2, "The final generated answer with inline citations is returned to the user.", level=1, size=14)

add_level_text(tf2, "6. EVALUATION - Measure and Improve", level=0, bold=True, size=16)
add_level_text(tf2, "golden_set.json acts as ground truth tests for the RAG system.", level=1, size=14)
add_level_text(tf2, "eval_retrieval.py evaluates search accuracy (Hit@k, MRR, etc.).", level=1, size=14)
add_level_text(tf2, "eval_generation.py evaluates LLM answer quality (faithfulness).", level=1, size=14)
add_level_text(tf2, "metrics.py compiles reports to track improvements over time.", level=1, size=14)

prs.save('D:\\ATCS-CPE\\ATCS-CPE-main\\LAB03\\RAG-Project\\RAG_System_Workflow_Updated.pptx')
print("Successfully generated updated PPTX.")
